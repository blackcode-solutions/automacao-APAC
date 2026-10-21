import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from apac_bot.config import MODALIDADE_GLAUCOMA, MODALIDADE_PADRAO, AppConfig
from apac_bot.sheet import (
    SheetError,
    create_example_workbook,
    extrair_codigos_associados,
    listar_abas_faturamento_glaucoma,
    listar_lancamentos_a_partir,
    load_lancamentos,
)


@pytest.fixture
def cfg_tmp(tmp_path: Path) -> AppConfig:
    example = json.loads(
        Path("config.example.json").read_text(encoding="utf-8")
    )
    xlsx = tmp_path / "planilha.xlsx"
    create_example_workbook(xlsx)
    example["planilha"]["caminho"] = str(xlsx)
    cfg_path = tmp_path / "config.json"
    cfg_path.write_text(json.dumps(example), encoding="utf-8")
    return AppConfig(raw=example, path=cfg_path)


def test_preserva_zeros_e_numeracao(cfg_tmp: AppConfig):
    result = load_lancamentos(cfg_tmp)
    assert result.ok
    primeiro = result.lancamentos[0]
    assert primeiro.codigo_original == "0412345678"
    assert primeiro.numeracao_original == "1234567890123-4"
    assert primeiro.numeracao_apac == "1234567890123/4"
    assert primeiro.id_paciente == "BIO-1001"


def test_duplicidade_numeracao(cfg_tmp: AppConfig, tmp_path: Path):
    path = tmp_path / "dup.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append(["PACIENTE", "COD.", "AUDITOR", "NUMERAÇÃO", "QUANT.", "ID_PACIENTE"])
    ws.append(["A", "01", "AUD", "111-1", "", "X"])
    ws.append(["B", "02", "AUD", "111-1", "", "Y"])
    for row in (2, 3):
        for col in (2, 4, 6):
            ws.cell(row, col).number_format = "@"
    wb.save(path)
    cfg_tmp.raw["planilha"]["caminho"] = str(path)
    result = load_lancamentos(cfg_tmp)
    assert not result.ok
    assert any("duplicada" in e.lower() for e in result.erros)


def test_quant_lido_sem_interpretacao(cfg_tmp: AppConfig):
    result = load_lancamentos(cfg_tmp)
    assert result.lancamentos[0].quant == "1"
    # segundo registro sem QUANT.
    assert result.lancamentos[1].quant is None


def test_linha_cabecalho_customizada(tmp_path: Path):
    path = tmp_path / "cab.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Prod"
    ws.append(["titulo"])
    ws.append([])
    ws.append(["QUANT.", "PACIENTE", "COD.", "AUDITOR", "NUMERAÇÃO"])
    ws.append(["1", "TESTE", "04123", "AUD", "111-1"])
    for col in (3, 5):
        ws.cell(4, col).number_format = "@"
        ws.cell(4, col).value = str(ws.cell(4, col).value)
    wb.save(path)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"]["aba"] = "Prod"
    example["planilha"]["linha_cabecalho"] = 3
    example["planilha"]["colunas"] = {
        "PACIENTE": "PACIENTE",
        "COD": "COD.",
        "AUDITOR": "AUDITOR",
        "NUMERACAO": "NUMERAÇÃO",
        "QUANT": "QUANT.",
        "ID_PACIENTE": "ID_PACIENTE",
    }
    cfg_path = tmp_path / "c.json"
    cfg_path.write_text(json.dumps(example), encoding="utf-8")
    cfg = AppConfig(raw=example, path=cfg_path)
    result = load_lancamentos(cfg)
    assert result.ok
    assert len(result.lancamentos) == 1
    assert result.lancamentos[0].codigo_original == "04123"


def test_listar_lancamentos_a_partir_meio_da_fila(cfg_tmp: AppConfig):
    nomes = [l.paciente_nome for l in load_lancamentos(cfg_tmp).lancamentos]
    assert len(nomes) >= 3
    inicio = nomes[1]
    cfg_tmp.raw.setdefault("apac", {})["paciente_inicio"] = inicio
    fila = listar_lancamentos_a_partir(cfg_tmp)
    assert [l.paciente_nome for l in fila] == nomes[1:]
    assert len(fila) == len(nomes) - 1


def test_listar_lancamentos_a_partir_sem_paciente_inicio_retorna_todos(
    cfg_tmp: AppConfig,
):
    todos = load_lancamentos(cfg_tmp).lancamentos
    assert len(todos) >= 1
    cfg_tmp.raw.setdefault("apac", {})["paciente_inicio"] = ""
    fila = listar_lancamentos_a_partir(cfg_tmp)
    assert [l.paciente_nome for l in fila] == [l.paciente_nome for l in todos]
    assert fila[0].paciente_nome == todos[0].paciente_nome


def test_listar_lancamentos_a_partir_paciente_inexistente(cfg_tmp: AppConfig):
    cfg_tmp.raw.setdefault("apac", {})["paciente_inicio"] = "NOME QUE NAO EXISTE"
    with pytest.raises(SheetError, match="não encontrado"):
        listar_lancamentos_a_partir(cfg_tmp)


def test_apenas_ativos_ignora_linhas_ocultas(tmp_path: Path):
    path = tmp_path / "filtro.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append(["PACIENTE", "COD.", "AUDITOR", "NUMERAÇÃO", "QUANT.", "ID_PACIENTE"])
    ws.append(["OCULTO UM", "01", "AUD", "111-1", "1", ""])
    ws.append(["VISIVEL UM", "02", "AUD", "222-2", "1", ""])
    ws.append(["OCULTO DOIS", "01", "AUD", "333-3", "1", ""])
    ws.append(["VISIVEL DOIS", "02", "AUD", "444-4", "1", ""])
    for row in range(2, 6):
        for col in (2, 4):
            ws.cell(row, col).number_format = "@"
            ws.cell(row, col).value = str(ws.cell(row, col).value)
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[4].hidden = True
    wb.save(path)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"]["apenas_ativos"] = True
    cfg_path = tmp_path / "c.json"
    cfg_path.write_text(json.dumps(example), encoding="utf-8")
    cfg = AppConfig(raw=example, path=cfg_path)

    result = load_lancamentos(cfg)
    assert result.ok
    assert len(result.lancamentos) == 2
    assert [l.paciente_nome for l in result.lancamentos] == ["VISIVEL UM", "VISIVEL DOIS"]
    assert any("oculta" in a.lower() for a in result.avisos)

    example["planilha"]["apenas_ativos"] = False
    cfg_all = AppConfig(raw=example, path=cfg_path)
    result_all = load_lancamentos(cfg_all)
    assert result_all.ok
    assert len(result_all.lancamentos) == 4


def test_extrair_codigos_associados_entre_parenteses():
    assert extrair_codigos_associados("") == []
    assert extrair_codigos_associados(None) == []
    assert extrair_codigos_associados(
        "FOB (0303050160) COL (0303050187)"
    ) == ["0303050160", "0303050187"]
    assert extrair_codigos_associados("(0303050160)(0303050187)") == [
        "0303050160",
        "0303050187",
    ]


def _xlsx_cenario_glaucoma(
    path: Path,
    *,
    chave_trat: str | None = "212620199580-8",
    chave_consulta: str | None = "212620199680-9",
    associado_trat: str | None = None,
) -> None:
    """Layout NOVO: Base Consultas + TRAT * com Chave (Guia); Relatório sem Guia."""
    wb = Workbook()
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.append(["ignore"])

    ws_rel = wb.create_sheet("Relatório")
    ws_rel.append(
        [
            "Prontuário",
            "Nº Atendimento",
            "Cliente",
            "Data",
            "Hora",
            "Guia",
            "Cód./Procedimento",
        ]
    )
    ws_rel.append(
        ["1956670", "7", "FRANCISCO DAS CHAGAS ALVES DA SILVA", "", "", None, "0303050160"]
    )

    headers_fat = [
        "Prontuário",
        "Nº Atend.",
        "Cliente",
        "Data",
        "Código",
        "Procedimento",
        "Valor (Total Conta)",
        "Associado a",
        "Chave (Guia)",
    ]
    ws_base = wb.create_sheet("Base Consultas")
    ws_base.append(headers_fat)
    ws_base.append(
        [
            "1991165",
            "2",
            "ROSELIA DIAMANTINA PINHEIRO",
            "12/08/2026",
            "0301010102",
            "Consulta",
            "57.74",
            "",
            chave_consulta,
        ]
    )

    ws_trat = wb.create_sheet("TRAT 0303050160")
    ws_trat.append(headers_fat)
    ws_trat.append(
        [
            "1956670",
            "7",
            "FRANCISCO DAS CHAGAS ALVES DA SILVA",
            "11/08/2026",
            "0303050160",
            "Tratamento",
            "98.04",
            associado_trat or "",
            chave_trat,
        ]
    )
    for ws in (ws_base, ws_trat):
        for row in range(2, 3):
            for col in (1, 2, 5, 9):
                cell = ws.cell(row, col)
                if cell.value is not None:
                    cell.number_format = "@"
                    cell.value = str(cell.value)
    wb.save(path)


def test_glaucoma_le_so_aba_selecionada(tmp_path: Path):
    path = tmp_path / "cenario.xlsx"
    _xlsx_cenario_glaucoma(path)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"]["aba"] = "PRODUÇÃO APAC SETEMBRO 2026"
    example["planilha"]["aba_glaucoma"] = "TRAT 0303050160"
    example.setdefault("apac", {})["modalidade"] = MODALIDADE_GLAUCOMA
    cfg = AppConfig(raw=example, path=tmp_path / "c.json")

    result = load_lancamentos(cfg)
    assert result.ok, result.erros
    assert len(result.lancamentos) == 1
    lac = result.lancamentos[0]
    assert lac.codigo_original == "0303050160"
    assert lac.prontuario == "1956670"
    assert lac.paciente_nome.startswith("FRANCISCO")
    assert lac.numeracao_original == "212620199580-8"
    assert lac.codigos_associados == []


def test_glaucoma_le_codigos_associados(tmp_path: Path):
    path = tmp_path / "cenario.xlsx"
    _xlsx_cenario_glaucoma(
        path,
        associado_trat="FOB (0303050187) COL (0303050195)",
    )

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"]["aba_glaucoma"] = "TRAT 0303050160"
    example.setdefault("apac", {})["modalidade"] = MODALIDADE_GLAUCOMA
    cfg = AppConfig(raw=example, path=tmp_path / "c.json")

    result = load_lancamentos(cfg)
    assert result.ok, result.erros
    lac = result.lancamentos[0]
    assert lac.codigo_apac == "0303050160"
    assert lac.codigos_associados == ["0303050187", "0303050195"]


def test_glaucoma_sem_aba_e_multiplas_faturamento_erro(tmp_path: Path):
    path = tmp_path / "cenario.xlsx"
    _xlsx_cenario_glaucoma(path)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"].pop("aba_glaucoma", None)
    example.setdefault("apac", {})["modalidade"] = MODALIDADE_GLAUCOMA
    cfg = AppConfig(raw=example, path=tmp_path / "c.json")

    with pytest.raises(SheetError, match="Selecione a aba"):
        load_lancamentos(cfg)


def test_glaucoma_uma_aba_faturamento_automatica(tmp_path: Path):
    path = tmp_path / "so_trat.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "TRAT 0303050187"
    ws.append(
        [
            "Prontuário",
            "Nº Atend.",
            "Cliente",
            "Data",
            "Código",
            "Procedimento",
            "Valor (Total Conta)",
            "Chave (Guia)",
        ]
    )
    ws.append(
        [
            "100",
            "1",
            "PACIENTE UNICA",
            "01/01/2026",
            "0303050187",
            "Trat",
            "10",
            "999-1",
        ]
    )
    wb.save(path)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example.setdefault("apac", {})["modalidade"] = MODALIDADE_GLAUCOMA
    cfg = AppConfig(raw=example, path=tmp_path / "c.json")

    result = load_lancamentos(cfg)
    assert result.ok, result.erros
    assert len(result.lancamentos) == 1
    assert result.lancamentos[0].numeracao_original == "999-1"


def test_listar_abas_faturamento_glaucoma(tmp_path: Path):
    path = tmp_path / "cenario.xlsx"
    _xlsx_cenario_glaucoma(path)
    abas = listar_abas_faturamento_glaucoma(path)
    assert abas == ["Base Consultas", "TRAT 0303050160"]


def test_glaucoma_sem_chave_guia_nao_gera_lancamentos(tmp_path: Path):
    path = tmp_path / "cenario.xlsx"
    _xlsx_cenario_glaucoma(path, chave_trat=None, chave_consulta=None)

    example = json.loads(Path("config.example.json").read_text(encoding="utf-8"))
    example["planilha"]["caminho"] = str(path)
    example["planilha"]["aba_glaucoma"] = "TRAT 0303050160"
    example.setdefault("apac", {})["modalidade"] = MODALIDADE_GLAUCOMA
    cfg = AppConfig(raw=example, path=tmp_path / "c.json")

    result = load_lancamentos(cfg)
    assert not result.ok
    assert any("sem lançamentos" in e.casefold() for e in result.erros)


def test_padrao_nao_usa_leitor_glaucoma(cfg_tmp: AppConfig):
    """Regressão: modalidade padrao continua no leitor oftalmo."""
    cfg_tmp.raw.setdefault("apac", {})["modalidade"] = MODALIDADE_PADRAO
    result = load_lancamentos(cfg_tmp)
    assert result.ok
    assert result.lancamentos[0].codigo_original == "0412345678"
    assert result.lancamentos[0].prontuario is None or result.lancamentos[0].prontuario == ""
