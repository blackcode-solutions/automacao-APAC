from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook, load_workbook

from apac_bot.config import MODALIDADE_GLAUCOMA, AppConfig, config_dir, resolve_path
from apac_bot.models import Lancamento
from apac_bot.normalize import cell_as_text, normalize_codigo, normalize_numeracao, optional_text

log = logging.getLogger(__name__)


class SheetError(Exception):
    pass


def _sem_acento(texto: str) -> str:
    t = unicodedata.normalize("NFKD", texto or "")
    return "".join(ch for ch in t if not unicodedata.combining(ch))


def nome_chave_prontuario(nome: str) -> str:
    """Chave estavel para cruzar PACIENTE <-> Cliente (casefold, espacos colapsados)."""
    return " ".join(_sem_acento(nome).casefold().split())


def prontuario_so_digitos(valor: Any) -> str:
    return "".join(ch for ch in cell_as_text(valor) if ch.isdigit())


def opcao_contem_prontuario(texto_opcao: str, prontuario: str) -> bool:
    """True se o texto da opcao do dropdown contem o prontuario como token numerico."""
    alvo = prontuario_so_digitos(prontuario)
    if not alvo:
        return False
    digitos = "".join(ch for ch in (texto_opcao or "") if ch.isdigit() or ch.isspace())
    # Tambem busca no texto bruto (ID pode aparecer colado ao nome)
    if re.search(rf"(?<!\d){re.escape(alvo)}(?!\d)", texto_opcao or ""):
        return True
    return bool(re.search(rf"(?<!\d){re.escape(alvo)}(?!\d)", digitos))


def sanitizar_nome_opcao(
    texto_opcao: str,
    prontuario: Optional[str] = None,
) -> str:
    """
    Extrai o nome legivel do texto da opcao do dropdown BIODATA.
    Remove prefixo de prontuario/ID. Retorna '' se sobrar so digitos.
    """
    t = " ".join((texto_opcao or "").replace("\n", " ").split())
    if not t:
        return ""
    pront = prontuario_so_digitos(prontuario or "")
    if pront:
        # Remove o prontuario como token (inicio ou no meio)
        t2 = re.sub(rf"(?<!\d){re.escape(pront)}(?!\d)", " ", t)
        t2 = " ".join(t2.split())
        if t2 and not t2.isdigit():
            t = t2
    # Prefixo numerico generico: "1763812 SANDRA MARIA" → "SANDRA MARIA"
    m = re.match(r"^\d{4,}\s+(.+)$", t)
    if m:
        t = m.group(1).strip()
    # So digitos = prontuario puro — invalido como nome
    if not t or t.isdigit() or re.fullmatch(r"[\d\s.\-/]+", t):
        return ""
    return t


def nome_mae_valido(nome_mae: Optional[str], paciente_nome: Optional[str]) -> bool:
    """False se mae vazia ou igual ao nome do paciente (contaminacao)."""
    mae = " ".join((nome_mae or "").split())
    if not mae:
        return False
    pac = " ".join((paciente_nome or "").split())
    if pac and nome_chave_prontuario(mae) == nome_chave_prontuario(pac):
        return False
    if mae.isdigit():
        return False
    return True


@dataclass
class SheetValidation:
    lancamentos: list[Lancamento]
    erros: list[str]
    avisos: list[str]

    @property
    def ok(self) -> bool:
        return not self.erros


def carregar_mapa_prontuario(cfg: AppConfig) -> dict[str, str]:
    """
    Le planilha_prontuarios (Cliente -> Prontuario).
    Retorna dict vazio se nao configurado ou arquivo ausente.
    """
    bloco = cfg.planilha_prontuarios
    if not bloco or not bloco.get("caminho"):
        return {}

    caminho = resolve_path(config_dir(cfg), Path(str(bloco["caminho"])))
    if not caminho.is_file():
        log.warning("planilha_prontuarios nao encontrada: %s", caminho)
        return {}

    aba = str(bloco.get("aba") or "APAC ").strip() or "APAC "
    # Aba no Excel pode ter espaco final ("APAC ")
    aba_cfg = str(bloco.get("aba") or "APAC ")
    try:
        header_row = int(bloco.get("linha_cabecalho") or 3)
    except (TypeError, ValueError):
        header_row = 3
    colunas = dict(bloco.get("colunas") or {})
    rotulo_pront = str(colunas.get("PRONTUARIO") or "Prontuário")
    rotulo_cli = str(colunas.get("CLIENTE") or "Cliente")

    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = None
        for nome in wb.sheetnames:
            if nome == aba_cfg or nome.strip() == aba or nome.casefold() == aba_cfg.casefold():
                ws = wb[nome]
                break
        if ws is None:
            log.warning(
                "Aba prontuarios %r nao encontrada em %s (abas=%s)",
                aba_cfg,
                caminho.name,
                wb.sheetnames,
            )
            return {}

        headers_raw = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i == header_row:
                headers_raw = list(row or ())
                break
        if not headers_raw:
            log.warning("Cabecalho prontuarios vazio (linha %s)", header_row)
            return {}

        def _hdr_key(h: Any) -> str:
            return nome_chave_prontuario(cell_as_text(h))

        alvo_pront = _hdr_key(rotulo_pront)
        alvo_cli = _hdr_key(rotulo_cli)
        idx_pront = idx_cli = None
        for j, h in enumerate(headers_raw):
            k = _hdr_key(h)
            if k == alvo_pront or k.startswith("prontuario"):
                idx_pront = j
            if k == alvo_cli or k == "cliente":
                idx_cli = j
        if idx_pront is None or idx_cli is None:
            log.warning(
                "Colunas Prontuario/Cliente nao achadas no cabecalho %s",
                [cell_as_text(h) for h in headers_raw],
            )
            return {}

        mapa: dict[str, str] = {}
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            if not row:
                continue
            pront = prontuario_so_digitos(row[idx_pront] if idx_pront < len(row) else "")
            nome = cell_as_text(row[idx_cli] if idx_cli < len(row) else "").strip()
            if not pront or not nome:
                continue
            chave = nome_chave_prontuario(nome)
            if chave in mapa and mapa[chave] != pront:
                log.warning(
                    "Prontuario duplicado para %r: mantendo %s (ignorando %s)",
                    nome,
                    mapa[chave],
                    pront,
                )
                continue
            if chave not in mapa:
                mapa[chave] = pront
        log.info("Mapa prontuarios: %s nome(s) carregado(s) de %s", len(mapa), caminho.name)
        return mapa
    finally:
        wb.close()


def enriquecer_prontuario(lanc: Lancamento, mapa: dict[str, str]) -> Lancamento:
    """Preenche lanc.prontuario pelo nome se ainda vazio."""
    if lanc.prontuario:
        return lanc
    if not mapa:
        return lanc
    chave = nome_chave_prontuario(lanc.paciente_nome or "")
    pront = mapa.get(chave)
    if pront:
        lanc.prontuario = pront
    return lanc


def listar_lancamentos_a_partir(
    cfg: AppConfig,
    planilha_override: Optional[Path] = None,
) -> list[Lancamento]:
    """
    Lançamentos ativos a partir de apac.paciente_inicio (nome exato, case-insensitive),
    inclusive até o fim da planilha.

    Se paciente_inicio estiver vazio/ausente, retorna todos os lançamentos ativos
    (partida no primeiro paciente com nome válido da planilha).
    """
    nome_alvo = (cfg.apac.get("paciente_inicio") or "").strip()

    result = load_lancamentos(cfg, planilha_override)
    if not result.ok:
        raise SheetError(
            "Falha ao carregar planilha: " + "; ".join(result.erros[:5])
        )
    if not result.lancamentos:
        raise SheetError("Planilha sem lançamentos ativos.")

    if not nome_alvo:
        return list(result.lancamentos)

    alvo_norm = nome_alvo.casefold()
    idx: Optional[int] = None
    for i, lanc in enumerate(result.lancamentos):
        if (lanc.paciente_nome or "").strip().casefold() == alvo_norm:
            idx = i
            break

    if idx is None:
        raise SheetError(
            f"Paciente início '{nome_alvo}' não encontrado na planilha "
            f"({len(result.lancamentos)} lançamento(s) ativos)."
        )
    return result.lancamentos[idx:]


def resolver_lancamento_inicio(
    cfg: AppConfig,
    planilha_override: Optional[Path] = None,
) -> Lancamento:
    """
    Localiza o lançamento de partida: por apac.paciente_inicio se configurado;
    senão o primeiro lançamento ativo da planilha.
    """
    return listar_lancamentos_a_partir(cfg, planilha_override)[0]


def _header_map(ws, header_row: int = 1) -> dict[str, int]:
    headers: dict[str, int] = {}
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        raw = ws.cell(header_row, col).value
        if raw is None:
            continue
        name = cell_as_text(raw)
        if name:
            headers[name] = col
    return headers


def _cell(ws, row: int, col: Optional[int]) -> str:
    if col is None:
        return ""
    return cell_as_text(ws.cell(row, col).value)


def _linha_oculta(ws, row: int) -> bool:
    """True se a linha estiver oculta (filtro do Excel / hide manual)."""
    try:
        return bool(ws.row_dimensions[row].hidden)
    except Exception:
        return False


ABA_GLAUCOMA_BASE_CONSULTAS = "Base Consultas"
ABA_GLAUCOMA_TRAT_PREFIX = "trat "

# Abas de faturamento do CENARIO (já trazem Chave/Guia). Relatório bruto não.
_GLAUCOMA_COLUNAS: dict[str, tuple[str, ...]] = {
    "prontuario": ("Prontuário", "Prontuario"),
    "cliente": ("Cliente",),
    "codigo": (
        "Código",
        "Codigo",
        "Cód./Procedimento",
        "Cod./Procedimento",
    ),
    "associado": ("Associado a", "Associados a", "Associado"),
    "guia": ("Chave (Guia)", "Guia"),
    "atendimento": (
        "Nº Atend.",
        "N° Atend.",
        "No Atend.",
        "Nº Atendimento",
        "N° Atendimento",
        "No Atendimento",
        "Num Atendimento",
    ),
}

_RE_CODIGO_ENTRE_PARENTESES = re.compile(r"\((\d+)\)")


def extrair_codigos_associados(texto: object) -> list[str]:
    """Extrai códigos de procedimento entre parênteses na coluna Associado a."""
    return _RE_CODIGO_ENTRE_PARENTESES.findall(str(texto or ""))


def _abas_faturamento_glaucoma(sheetnames: list[str]) -> list[str]:
    """Base Consultas + abas TRAT * (ordem do arquivo)."""
    abas: list[str] = []
    for nome in sheetnames:
        chave = nome.strip().casefold()
        if chave == ABA_GLAUCOMA_BASE_CONSULTAS.casefold():
            abas.append(nome)
        elif chave.startswith(ABA_GLAUCOMA_TRAT_PREFIX):
            abas.append(nome)
    return abas


def listar_abas_faturamento_glaucoma(caminho: Path) -> list[str]:
    """Abas de faturamento do xlsx glaucoma (para o dropdown da UI)."""
    path = Path(caminho)
    if not path.is_file():
        raise SheetError(f"Planilha não encontrada: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return _abas_faturamento_glaucoma(list(wb.sheetnames))
    finally:
        wb.close()


def _resolver_abas_glaucoma(
    abas: list[str], _aba_escolhida: str = ""
) -> list[str]:
    """
    Todas as abas de faturamento descobertas no xlsx.

    _aba_escolhida é ignorado (legado); o Iniciar sempre processa a lista completa.
    """
    return list(abas)


def _indice_coluna_glaucoma(
    headers: dict[str, int], aliases: tuple[str, ...]
) -> Optional[int]:
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    lower = {k.casefold(): v for k, v in headers.items()}
    for alias in aliases:
        hit = lower.get(alias.casefold())
        if hit is not None:
            return hit
    return None


def _ler_aba_faturamento_glaucoma(
    ws,
    *,
    nome_aba: str,
    cfg: AppConfig,
    apenas_ativos: bool,
    vistas: dict[str, str],
) -> tuple[list[Lancamento], list[str], list[str], int]:
    """Lê uma aba Base Consultas / TRAT. Retorna (lancs, erros, avisos, ignorados)."""
    headers = _header_map(ws, header_row=1)
    erros: list[str] = []
    avisos: list[str] = []
    indices: dict[str, Optional[int]] = {}
    obrigatorias = ("cliente", "codigo", "prontuario", "guia")
    for chave in _GLAUCOMA_COLUNAS:
        indices[chave] = _indice_coluna_glaucoma(headers, _GLAUCOMA_COLUNAS[chave])
    for chave in obrigatorias:
        if indices[chave] is None:
            erros.append(
                f"Aba '{nome_aba}': coluna obrigatória ausente "
                f"'{_GLAUCOMA_COLUNAS[chave][0]}'"
            )
    if erros:
        return [], erros, avisos, 0

    regras_num = (cfg.normalizacao.get("numeracao") or {})
    regras_cod = (cfg.normalizacao.get("codigo") or {})
    fixos = cfg.dados_fixos
    lancamentos: list[Lancamento] = []
    ignorados = 0
    max_row = ws.max_row or 2

    for row in range(2, max_row + 1):
        paciente = _cell(ws, row, indices["cliente"])
        codigo = _cell(ws, row, indices["codigo"])
        prontuario = _cell(ws, row, indices["prontuario"])
        guia = _cell(ws, row, indices["guia"])
        n_atend = _cell(ws, row, indices.get("atendimento"))
        associado_bruto = _cell(ws, row, indices.get("associado"))

        if not any([paciente, codigo, prontuario, guia, n_atend]):
            continue
        # Sem Chave (Guia) não entra no lote (linhas residuais / incompletas).
        if not guia:
            continue
        if apenas_ativos and _linha_oculta(ws, row):
            ignorados += 1
            continue

        faltando = []
        if not paciente:
            faltando.append("Cliente")
        if not codigo:
            faltando.append(_GLAUCOMA_COLUNAS["codigo"][0])
        if not prontuario:
            faltando.append("Prontuário")
        if faltando:
            erros.append(
                f"Aba '{nome_aba}' linha {row}: campos ausentes: {', '.join(faltando)}"
            )
            continue

        numeracao = guia
        id_vista = f"{numeracao}|{codigo}"
        if id_vista in vistas:
            erros.append(
                f"Lançamento duplicado Chave (Guia) '{numeracao}' + código "
                f"'{codigo}' (abas/linhas {vistas[id_vista]} e {nome_aba}:{row})"
            )
            continue
        vistas[id_vista] = f"{nome_aba}:{row}"

        associados = [
            normalize_codigo(c, regras_cod)
            for c in extrair_codigos_associados(associado_bruto)
        ]

        lancamentos.append(
            Lancamento(
                clinica_id=cfg.clinica_id,
                numeracao_original=numeracao,
                numeracao_apac=normalize_numeracao(numeracao, regras_num),
                codigo_original=codigo,
                codigo_apac=normalize_codigo(codigo, regras_cod),
                auditor="",
                paciente_nome=paciente,
                prontuario=optional_text(prontuario),
                quant="1",
                competencia=optional_text(fixos.get("competencia")),
                validade=optional_text(fixos.get("validade")),
                linha_planilha=row,
                codigos_associados=associados,
            )
        )

    return lancamentos, erros, avisos, ignorados


def _load_lancamentos_glaucoma(
    cfg: AppConfig, planilha_override: Optional[Path] = None
) -> SheetValidation:
    """
    Leitor do CENARIO glaucoma: todas as abas Base Consultas / TRAT *
    descobertas no xlsx (nomes variam por arquivo). Número da APAC = Chave (Guia).
    """
    base = config_dir(cfg)
    caminho = planilha_override or resolve_path(base, cfg.planilha_caminho)
    if not caminho.is_file():
        raise SheetError(f"Planilha não encontrada: {caminho}")

    apenas_ativos = cfg.planilha_apenas_ativos
    wb = load_workbook(caminho, data_only=True, read_only=not apenas_ativos)
    try:
        abas = _abas_faturamento_glaucoma(list(wb.sheetnames))
        if not abas:
            raise SheetError(
                "Nenhuma aba de faturamento glaucoma encontrada "
                f"(esperado 'Base Consultas' e/ou 'TRAT …'). "
                f"Disponíveis: {wb.sheetnames}"
            )
        abas = _resolver_abas_glaucoma(abas)

        erros: list[str] = []
        avisos: list[str] = [
            f"Glaucoma: {len(abas)} aba(s): {', '.join(abas)}"
        ]
        lancamentos: list[Lancamento] = []
        vistas: dict[str, str] = {}
        ignorados_filtro = 0

        for nome_aba in abas:
            ws = wb[nome_aba]
            lacs, errs, avis, ign = _ler_aba_faturamento_glaucoma(
                ws,
                nome_aba=nome_aba,
                cfg=cfg,
                apenas_ativos=apenas_ativos,
                vistas=vistas,
            )
            lancamentos.extend(lacs)
            erros.extend(errs)
            avisos.extend(avis)
            ignorados_filtro += ign

        if apenas_ativos and ignorados_filtro:
            avisos.append(
                f"Apenas ativos: {len(lancamentos)} lançamento(s) visível(is); "
                f"{ignorados_filtro} linha(s) oculta(s) ignorada(s) pelo filtro do Excel."
            )
        if not lancamentos and not erros:
            erros.append(
                "Planilha glaucoma sem lançamentos nas abas de faturamento "
                f"({', '.join(abas)})."
            )
        return SheetValidation(lancamentos, erros, avisos)
    finally:
        wb.close()


def load_lancamentos(
    cfg: AppConfig, planilha_override: Optional[Path] = None
) -> SheetValidation:
    if cfg.modalidade == MODALIDADE_GLAUCOMA:
        return _load_lancamentos_glaucoma(cfg, planilha_override)

    base = config_dir(cfg)
    caminho = planilha_override or resolve_path(base, cfg.planilha_caminho)
    if not caminho.is_file():
        raise SheetError(f"Planilha não encontrada: {caminho}")

    apenas_ativos = cfg.planilha_apenas_ativos
    # read_only=True não expõe bem row_dimensions.hidden; use modo normal com apenas_ativos
    wb = load_workbook(caminho, data_only=True, read_only=not apenas_ativos)
    aba = cfg.planilha_aba
    if aba not in wb.sheetnames:
        raise SheetError(f"Aba '{aba}' não encontrada. Disponíveis: {wb.sheetnames}")
    ws = wb[aba]

    colmap = cfg.colunas
    header_row = cfg.planilha_linha_cabecalho
    headers = _header_map(ws, header_row)
    erros: list[str] = []
    avisos: list[str] = []

    required_headers = {
        "PACIENTE": colmap["PACIENTE"],
        "COD": colmap["COD"],
        "AUDITOR": colmap["AUDITOR"],
        "NUMERACAO": colmap["NUMERACAO"],
    }
    indices: dict[str, Optional[int]] = {}
    for key, header_name in required_headers.items():
        if header_name not in headers:
            erros.append(f"Coluna obrigatória ausente na planilha: '{header_name}'")
            indices[key] = None
        else:
            indices[key] = headers[header_name]

    for opt_key in ("QUANT", "ID_PACIENTE"):
        header_name = colmap.get(opt_key)
        if header_name and header_name in headers:
            indices[opt_key] = headers[header_name]
        else:
            indices[opt_key] = None
            if opt_key == "ID_PACIENTE" and header_name:
                avisos.append(
                    f"Coluna opcional '{header_name}' não encontrada; busca priorizará nome."
                )

    if erros:
        wb.close()
        return SheetValidation([], erros, avisos)

    regras_num = (cfg.normalizacao.get("numeracao") or {})
    regras_cod = (cfg.normalizacao.get("codigo") or {})
    fixos = cfg.dados_fixos

    lancamentos: list[Lancamento] = []
    vistas: dict[str, int] = {}
    ignorados_filtro = 0

    data_start = header_row + 1
    max_row = ws.max_row or data_start
    for row in range(data_start, max_row + 1):
        paciente = _cell(ws, row, indices["PACIENTE"])
        codigo = _cell(ws, row, indices["COD"])
        auditor = _cell(ws, row, indices["AUDITOR"])
        numeracao = _cell(ws, row, indices["NUMERACAO"])
        quant = _cell(ws, row, indices.get("QUANT"))
        id_paciente = _cell(ws, row, indices.get("ID_PACIENTE"))

        if not any([paciente, codigo, auditor, numeracao, quant, id_paciente]):
            continue

        if apenas_ativos and _linha_oculta(ws, row):
            ignorados_filtro += 1
            continue

        faltando = []
        if not paciente:
            faltando.append("PACIENTE")
        if not codigo:
            faltando.append("COD")
        if not auditor:
            faltando.append("AUDITOR")
        if not numeracao:
            faltando.append("NUMERACAO")
        if faltando:
            erros.append(f"Linha {row}: campos ausentes: {', '.join(faltando)}")
            continue

        if numeracao in vistas:
            erros.append(
                f"NUMERAÇÃO duplicada '{numeracao}' nas linhas {vistas[numeracao]} e {row}"
            )
            continue
        vistas[numeracao] = row

        # QUANT. é lido sem interpretação como quantidade de procedimentos
        lancamentos.append(
            Lancamento(
                clinica_id=cfg.clinica_id,
                numeracao_original=numeracao,
                numeracao_apac=normalize_numeracao(numeracao, regras_num),
                codigo_original=codigo,
                codigo_apac=normalize_codigo(codigo, regras_cod),
                auditor=auditor,
                paciente_nome=paciente,
                id_paciente=optional_text(id_paciente),
                quant=optional_text(quant),
                competencia=optional_text(fixos.get("competencia")),
                validade=optional_text(fixos.get("validade")),
                linha_planilha=row,
            )
        )

    if apenas_ativos and ignorados_filtro:
        avisos.append(
            f"Apenas ativos: {len(lancamentos)} lançamento(s) visível(is); "
            f"{ignorados_filtro} linha(s) oculta(s) ignorada(s) pelo filtro do Excel."
        )

    wb.close()

    # Glaucoma: prontuário já vem na planilha do lote — não usa 2ª planilha.
    if cfg.modalidade != MODALIDADE_GLAUCOMA:
        mapa = carregar_mapa_prontuario(cfg)
        if mapa:
            for lanc in lancamentos:
                enriquecer_prontuario(lanc, mapa)

    return SheetValidation(lancamentos, erros, avisos)


def create_example_workbook(path: Path) -> None:
    """Gera planilha fictícia para testes e exemplos."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    headers = ["PACIENTE", "COD.", "AUDITOR", "NUMERAÇÃO", "QUANT.", "ID_PACIENTE"]
    ws.append(headers)
    rows = [
        ["MARIA SILVA FICTICIA", "0412345678", "AUD01", "1234567890123-4", "1", "BIO-1001"],
        ["JOAO SOUZA FICTICIO", "0412345678", "AUD01", "1234567890124-5", "", "BIO-1002"],
        ["ANA COSTA FICTICIA", "0400000123", "AUD02", "1234567890125-6", "2", ""],
        ["PACIENTE AMBIGUO", "0412345678", "AUD01", "1234567890126-7", "", ""],
    ]
    for r in rows:
        ws.append(r)
    # Garantir que códigos/numerações com zero à esquerda sejam texto
    for row_idx in range(2, 2 + len(rows)):
        for col_idx in (2, 4, 6):  # COD., NUMERAÇÃO, ID_PACIENTE
            cell = ws.cell(row_idx, col_idx)
            if cell.value is not None and cell.value != "":
                cell.value = str(cell.value)
                cell.number_format = "@"
    wb.save(path)
